from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

OUT=str(Path(__file__).resolve().parent/'HimanshuDigitalTwin_Presentation.pptx')
P=Presentation();P.slide_width=Inches(13.333);P.slide_height=Inches(7.5)
B=P.slide_layouts[6]
NAVY='152C3C'; INK='172F3B'; TEAL='167E78'; MINT='A5DCCB'; PALE='F5F8F7'; WHITE='FFFFFF'; MUTED='667D88'; GREY='D9E4E6'; ORANGE='B46F4B'; LIGHT='E6F1EF'; HBLUE='254C5A'

def rgb(h): return RGBColor.from_string(h)
def rect(s,x,y,w,h,fill=WHITE,line=None,rad=0):
 shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill);shape.line.fill.background() if not line else None
 if line:shape.line.color.rgb=rgb(line);shape.line.width=Pt(1)
 if rad:
  try:shape.adjustments[0]=0.15
  except:pass
 return shape

def circle(s,x,y,r,fill=TEAL):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(r), Inches(r));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background();return sh

def tx(s,t,x,y,w,h,sz=18,color=INK,bold=False,align='left',font='Aptos',val=MSO_ANCHOR.MIDDLE):
 z=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));z.text_frame.clear();z.text_frame.word_wrap=True;z.text_frame.margin_left=Inches(0);z.text_frame.margin_right=Inches(0);z.text_frame.margin_top=Inches(0);z.text_frame.margin_bottom=Inches(0);z.text_frame.vertical_anchor=val
 for i,part in enumerate(t.split('\n')):
  p=z.text_frame.paragraphs[0] if i==0 else z.text_frame.add_paragraph();p.text=part;p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align];p.space_after=Pt(2)
  p.font.size=Pt(sz);p.font.bold=bold;p.font.name=font;p.font.color.rgb=rgb(color)
 return z

def line(s,x,y,xx,yy,c=GREY,weight=1.4):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(xx), Inches(yy));sh.line.color.rgb=rgb(c);sh.line.width=Pt(weight);return sh

def base(num,tag,title,subtitle=None,dark=False):
 s=P.slides.add_slide(B);rect(s,0,0,13.333,7.5,NAVY if dark else PALE);rect(s,0,0,.09,7.5,TEAL)
 if tag: tx(s,tag.upper(),.6,.34,10,.27,10,MINT if dark else TEAL,True)
 if title: tx(s,title,.6,.75,11.95,1.12,34,WHITE if dark else INK,True,val=MSO_ANCHOR.MIDDLE)
 if subtitle:tx(s,subtitle,.62,1.86,11.9,.45,15,'BED3D8' if dark else MUTED)
 line(s,.6,7.05,12.75,7.05,'395968' if dark else GREY,1)
 tx(s,'HIMANSHUDIGITALTWIN  /  MITS GWALIOR',.6,7.10,7,.2,9,'BED3D8' if dark else MUTED,True)
 tx(s,f'{num:02d} / 06',11.6,7.10,1.16,.2,9,'BED3D8' if dark else MUTED,True,align='right')
 return s

def card(s,x,y,w,h,fill=WHITE,border=None):return rect(s,x,y,w,h,fill,border,rad=.13)

# 01 — Cover: deliberately not generic stock illustration
s=base(1,'Happiest Health Digital Twin Challenge 2026',None,dark=True)
tx(s,'Himanshu',.62,1.1,7.1,.76,47,WHITE,True);tx(s,'DigitalTwin',.62,1.88,7.1,.82,50,MINT,True)
tx(s,'A research prototype for near-term\nglucose forecasting — with a reason\nto withhold a prediction.',.65,2.99,6.5,1.55,25,WHITE)
tx(s,'Team  himanshusubh1510',.65,5.5,5.9,.37,16,MINT,True)
tx(s,'Presenter  Himanshu Shukla  ·  MITS Gwalior',.65,5.92,6.8,.34,14,'D8E7E7')
card(s,7.7,1.45,4.95,4.72,HBLUE)
tx(s,'SYNTHETIC REPLAY  /  EXAMPLE',8.04,1.77,4.3,.27,11,MINT,True)
tx(s,'144',8.05,2.17,2,.91,58,WHITE,True);tx(s,'mg/dL  at 12:15',10.09,2.54,2.25,.37,14,'C1D6DB')
line(s,8.05,4.55,12.24,4.55,'658590',1);line(s,8.05,5.12,12.24,5.12,'658590',1)
ps=[(8.08,4.91),(9.35,4.36),(10.58,4.34),(12.16,4.82)];
for (a,b),(c,d) in zip(ps,ps[1:]):line(s,a,b,c,d,MINT,3)
for a,b in ps:circle(s,a-.06,b-.06,.12,MINT)
tx(s,'NOW',8.05,5.35,.8,.32,11,'B8CED1',True);tx(s,'+30',9.26,5.35,.9,.32,11,'B8CED1',True);tx(s,'+60',10.45,5.35,.9,.32,11,'B8CED1',True);tx(s,'+120',11.72,5.35,.9,.32,11,'B8CED1',True)
tx(s,'Illustrative numbers. Not a calibrated clinical forecast.',8.04,6.02,4.2,.32,10,'B8CED1')

# 02 — The clinical framing without overstating
s=base(2,'01 / The problem','One number describes the present. It does not explain the next two hours.','Patient context and measurement quality change what a forecast can mean.')
card(s,.65,2.42,4.0,3.67,WHITE)
tx(s,'AT 12:15 PM',.94,2.72,2.8,.33,12,TEAL,True);tx(s,'144',.95,3.22,3.03,1.07,63,INK,True);tx(s,'mg/dL  ·  synthetic observed value',.96,4.32,3.4,.45,15,MUTED)
line(s,.96,5.1,4.35,5.1,GREY)
tx(s,'A reading is evidence —\nnot a two-hour prediction.',.96,5.21,3.48,.62,15,INK,True)
card(s,4.87,2.42,7.78,3.67,WHITE)
tx(s,'ILLUSTRATIVE FORECAST FROM THE SAME REPLAY',5.22,2.73,7,.34,12,TEAL,True)
vals=[('NOW','144',.0),('+30 min','155.5',1),('+60 min','156.0',2),('+120 min','145.4',3)]
for i,(a,b,_) in enumerate(vals):
 x=5.22+i*1.77
 tx(s,a,x,3.29,1.58,.32,13,MUTED,True)
 tx(s,b,x,3.70,1.65,.65,27,INK,True)
 if i<3:line(s,x+1.51,4.04,x+1.7,4.04,TEAL,2)
line(s,5.22,4.83,12.24,4.83,GREY)
tx(s,'Context: 59.7 g recent carbohydrates  ·  heart rate 82 bpm  ·  time aligned CGM',5.22,5.03,7,.61,14,MUTED)
tx(s,'Question for the prototype: can we show useful context without hiding sensor uncertainty?',.69,6.37,11.88,.35,16,INK,True)

# 03 — architecture connections and safety gate
s=base(3,'02 / Architecture','Synchronize first. Predict only when supported.','The judge can trace an observation through an explicit quality decision.')
for i,(name,detail) in enumerate([('Patient profile','Age · BMI · HbA1c'),('CGM observations','Value · timestamp'),('Meal + activity','Carbs · movement')]):
 x=.65+i*2.46;card(s,x,2.50,2.27,1.10,WHITE,GREY);tx(s,name,x+.15,2.69,1.98,.35,15,INK,True);tx(s,detail,x+.15,3.1,2.0,.23,11,MUTED)
card(s,1.74,4.05,4.05,1.02,LIGHT)
tx(s,'Versioned patient state',1.96,4.26,3.65,.31,18,INK,True);tx(s,'Timestamp alignment + replay history',1.96,4.66,3.65,.25,11,MUTED)
for x in [1.77,4.23,6.68]:line(s,x,3.61,x,3.81,TEAL,1.6)
line(s,1.77,3.81,4.23,3.81,TEAL);line(s,4.23,3.81,4.23,4.05,TEAL)
line(s,5.79,4.57,6.45,4.57,TEAL,2)
card(s,6.45,4.02,2.43,1.10,'E8F2F1',TEAL)
tx(s,'DATA QUALITY',6.65,4.18,2.0,.28,11,TEAL,True);tx(s,'Fresh enough?',6.65,4.54,2.06,.34,17,INK,True)
line(s,8.89,4.33,9.43,3.26,TEAL,2);line(s,8.89,4.81,9.43,5.62,ORANGE,2)
card(s,9.44,2.76,3.08,1.08,'E4F1EC');tx(s,'VALID / SUPPORTED',9.62,2.9,2.8,.24,11,TEAL,True);tx(s,'30 / 60 / 120 min',9.62,3.23,2.8,.31,17,INK,True)
card(s,9.44,5.13,3.08,1.08,'F6EAE3');tx(s,'SENSOR TOO STALE',9.62,5.26,2.8,.25,11,ORANGE,True);tx(s,'Withhold forecast + risk',9.62,5.58,2.79,.33,15,INK,True)
line(s,9.4,6.46,2.39,6.46,'A6BCC2',1);line(s,2.39,6.46,2.39,5.08,'A6BCC2',1)
tx(s,'Later valid observations are compared with earlier forecasts (synthetic verification)',2.74,6.20,6.42,.22,11,MUTED)
tx(s,'Python · FastAPI · HTML/CSS/JS  /  scikit-learn research pipeline is separate',.66,6.73,11.2,.24,11,MUTED)

# 04 — genuine local dashboard screenshot with editable judge notes
s=base(4,'03 / Working prototype','The running demo — not just an architecture diagram.','Actual FastAPI dashboard, reproducible synthetic lunch checkpoint.')
shot=Path(__file__).resolve().parent/'Live_Demo_Lunch.png'
if shot.is_file():
  # Only embed a real image captured from a running localhost application.
  # Every annotation, heading and chapter list remains an editable PowerPoint object.
  card(s,.62,2.43,8.54,4.22,WHITE,GREY)
  pic=s.shapes.add_picture(str(shot), Inches(.74), Inches(2.55), width=Inches(8.29), height=Inches(3.75))
  tx(s,'ACTUAL LOCAL DEMO   /   LUNCH CHECKPOINT',.85,6.34,7.98,.22,10,TEAL,True)
  card(s,9.36,2.43,3.32,4.22,'E5F0EE')
  tx(s,'JUDGE PATH',9.63,2.69,2.75,.26,12,TEAL,True)
  for i,(name,detail) in enumerate([
     ('01 · Baseline','Initial reading'),('02 · Breakfast','Meal context'),
     ('03 · Lunch','Forecast trajectory'),('04 · Verification','Forecasts resolved'),
     ('05 · Sensor gap','Forecast withheld'),('06 · Recovery','Fresh signal returns')]):
      y=3.10+i*.49
      tx(s,name,9.63,y,2.8,.23,14,INK,True)
      tx(s,detail,9.63,y+.22,2.8,.2,10,MUTED)
  tx(s,'Six deterministic synthetic states. No clinical predictions.',9.61,6.17,2.78,.28,10,ORANGE)
else:
  # Intentionally transparent when no real capture has been obtained.
  card(s,.66,2.45,8.54,4.20,WHITE,GREY)
  tx(s,'Screenshot not yet captured',1.15,3.44,7.6,.65,28,INK,True)
  tx(s,'Run submission/capture_demo.py against the localhost app.\nNever substitute a fabricated dashboard image.',1.15,4.30,7.65,.87,18,MUTED)
  card(s,9.37,2.45,3.33,4.20,LIGHT)
  for i,label in enumerate(['BASELINE','BREAKFAST','LUNCH','VERIFICATION','SENSOR GAP','RECOVERY']):
    tx(s,f'{i+1:02d}  {label}',9.68,2.88+i*.50,2.75,.31,14,TEAL if i!=4 else ORANGE,True)
  tx(s,'Real application image required for final evidence.',.85,6.36,7.85,.28,12,ORANGE)

# 05 — distinguish transitions with 15-min rule
s=base(5,'04 / Sensor integrity','When the signal fails, the interface must say so.','The demo includes a deliberately missing/stale glucose segment.')
card(s,.67,2.41,5.72,3.36,'E7F3EF',None)
tx(s,'VALID OBSERVATION',1.02,2.75,4.8,.36,14,TEAL,True);tx(s,'144',1.03,3.35,3.05,.99,61,INK,True);tx(s,'mg/dL',3.98,3.75,1.5,.29,16,MUTED)
line(s,1.02,4.72,5.97,4.72,'BDD3CD')
tx(s,'Forecast available',1.03,4.89,4.7,.34,22,TEAL,True);tx(s,'Synthetic example at 12:15 PM',1.03,5.31,4.8,.26,12,MUTED)
card(s,6.6,2.41,6.05,3.36,'F7EDE7',None)
tx(s,'STALE SENSOR / 14:25',6.95,2.75,5,.37,14,ORANGE,True);tx(s,'WITHHELD',6.96,3.45,5.3,.81,38,INK,True)
line(s,6.97,4.72,12.17,4.72,'DDC7B8')
tx(s,'No verified glucose  ·  No forecast  ·  No risk',6.97,4.89,5.06,.75,18,ORANGE,True)
tx(s,'Up to 15 minutes: model may use the last valid reading for a provisional estimate.',.82,6.01,11.75,.4,15,INK)
tx(s,'Beyond the freshness rule: no estimate is shown. Imputed values are never drawn as measured CGM.',.82,6.45,11.75,.34,14,MUTED)

# 06 — separate research evidence and honest limits
s=base(6,'05 / Evidence and limits','A measured research signal. Not a deployed clinical model.','Independent CGMacros study, patient-grouped holdout; lower MAE is better.')
card(s,.65,2.36,7.45,3.94,WHITE)
tx(s,'MAE  /  mg/dL',.94,2.62,3,.35,12,TEAL,True)
for i,(label,base_val,model_val) in enumerate([('30 min',14.82,13.09),('60 min',22.94,20.77),('120 min',32.46,27.67)]):
 y=3.13+i*.91;tx(s,label,.96,y,1.15,.3,14,INK,True)
 scale=4.24/35
 rect(s,2.1,y,base_val*scale,.19,'9FB7C1')
 rect(s,2.1,y+.27,model_val*scale,.19,TEAL)
 tx(s,f'{base_val:.2f} / {model_val:.2f}',6.4,y-.01,1.37,.47,13,INK,True)
tx(s,'Gray: persistence baseline     Teal: independent study model',.96,5.95,6.89,.25,11,MUTED)
card(s,8.33,2.36,4.35,3.94,'E8F0F0')
tx(s,'EVIDENCE BOUNDARY',8.65,2.66,3.66,.27,12,TEAL,True)
tx(s,'45',8.65,3.08,1.12,.63,35,INK,True);tx(s,'participants  /  14 with T2D',9.77,3.26,2.63,.28,13,MUTED)
line(s,8.64,3.85,12.29,3.85,'C6D6D9')
tx(s,'Only 3 T2D participants in the\n10-person primary holdout.',8.65,4.08,3.55,.91,17,INK,True)
tx(s,'The study model is NOT the model\nrunning in the synthetic dashboard.',8.65,5.18,3.62,.80,15,ORANGE,True)
tx(s,'What the judge can verify: deterministic replay, sensor-gap abstention, and research provenance.',.73,6.50,12,.40,15,INK,True)

P.core_properties.title='HimanshuDigitalTwin — Happiest Health six-slide submission'
P.core_properties.subject='Research glucose digital twin: synthetic demo and independent evidence'
P.core_properties.author='Himanshu Shukla'
P.save(OUT)
print(OUT)